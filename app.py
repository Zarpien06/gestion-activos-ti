from collections import Counter
import io
import re
import time
import uuid
import zipfile
from datetime import datetime, timedelta, timezone
from pathlib import Path
import os

from dotenv import load_dotenv
from flask import (
    Flask,
    abort,
    flash,
    has_request_context,
    redirect,
    render_template,
    request,
    send_file,
    send_from_directory,
    session,
    url_for,
)
from flask_login import (
    current_user,
    login_required,
    login_user,
    logout_user,
)
from flask_wtf.csrf import CSRFError, CSRFProtect

from api.recolector_routes import recolector_bp
from auth import inicializar_login
from auth.auth_service import (
    actualizar_usuario,
    autenticar_usuario,
    cambiar_password,
    crear_usuario,
    limpiar_cache_usuarios,
)
from auth.decorators import permiso_requerido, solo_administrador
from services.pdf_generator import generar_acta_pdf
from supabase_client import supabase


load_dotenv()

app = Flask(__name__)
app.config["SECRET_KEY"] = os.getenv(
    "SECRET_KEY",
    "cambia-esta-clave-en-produccion",
)
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_SECURE"] = (
    os.getenv("RENDER", "").lower() == "true"
)
app.config["PERMANENT_SESSION_LIFETIME"] = timedelta(
    minutes=int(os.getenv("SESSION_TIMEOUT_MINUTES", "30"))
)

# Limite de tamano por envio (facturas del activo + adicionales)
app.config["MAX_CONTENT_LENGTH"] = 20 * 1024 * 1024  # 20 MB

# Cookie "Recordar sesion" (Flask-Login)
app.config["REMEMBER_COOKIE_DURATION"] = timedelta(days=7)
app.config["REMEMBER_COOKIE_HTTPONLY"] = True
app.config["REMEMBER_COOKIE_SAMESITE"] = "Lax"
app.config["REMEMBER_COOKIE_SECURE"] = (
    os.getenv("RENDER", "").lower() == "true"
)

inicializar_login(app)
app.register_blueprint(recolector_bp)

# Proteccion CSRF global. El recolector queda exento porque los equipos
# se autentican con token (tabla tokens_recolector), no con sesion.
csrf = CSRFProtect(app)
csrf.exempt(recolector_bp)


# El endpoint del recolector es publico para la sesion web porque
# los equipos se autentican con un token (tabla tokens_recolector),
# no con login.
RUTAS_PUBLICAS = {
    "login",
    "static",
    "recolector.recibir_equipo",
}


def cerrar_sesion_y_redirigir(mensaje=None, categoria="success", destino="login"):
    """Cierra la sesion de forma definitiva y manda al login.

    El orden importa: primero se limpia la sesion y despues se llama a
    logout_user(), para que Flask-Login pueda marcar la cookie de
    "recordar" para borrado. Ademas se borra la cookie de forma explicita.
    """
    session.clear()
    logout_user()

    if mensaje:
        # El flash va despues de limpiar la sesion, si no se perderia
        flash(mensaje, categoria)

    respuesta = redirect(url_for(destino))
    respuesta.delete_cookie(
        app.config.get("REMEMBER_COOKIE_NAME", "remember_token")
    )
    return respuesta


@app.before_request
def controlar_sesion():
    endpoint = request.endpoint

    if endpoint is None or endpoint in RUTAS_PUBLICAS:
        return None

    if not current_user.is_authenticated:
        flash(
            "Tu sesion expiro o debes iniciar sesion nuevamente.",
            "warning",
        )
        return redirect(url_for("login", next=request.path))

    ahora = datetime.now(timezone.utc)
    ultimo_acceso = session.get("ultimo_acceso")

    if ultimo_acceso:
        try:
            ultimo = datetime.fromisoformat(ultimo_acceso)
            limite = app.config["PERMANENT_SESSION_LIFETIME"]

            if ahora - ultimo > limite:
                nombre = current_user.nombre
                return cerrar_sesion_y_redirigir(
                    f"La sesion de {nombre} expiro por inactividad.",
                    "warning",
                )
        except (TypeError, ValueError):
            return cerrar_sesion_y_redirigir(
                "La sesion no es valida. Inicia sesion nuevamente.",
                "warning",
            )

    session["ultimo_acceso"] = ahora.isoformat()
    session.permanent = True
    return None


@app.after_request
def sin_cache(respuesta):
    """Evita que el navegador guarde paginas (boton 'atras' tras cerrar sesion)."""
    if request.endpoint != "static":
        respuesta.headers["Cache-Control"] = (
            "no-store, no-cache, must-revalidate, max-age=0"
        )
        respuesta.headers["Pragma"] = "no-cache"
        respuesta.headers["Expires"] = "0"
    return respuesta


# ==================================================
# FUNCIONES GENERALES
# ==================================================

def consultar_tabla(nombre_tabla, columnas="*", ordenar_por=None):
    # Reintenta una vez: Supabase a veces corta la conexion por inactividad
    for intento in range(2):
        try:
            consulta = supabase.table(nombre_tabla).select(columnas)
            if ordenar_por:
                consulta = consulta.order(ordenar_por)
            return consulta.execute().data or []
        except Exception as error:
            print(f"Error consultando {nombre_tabla} (intento {intento + 1}): {error}")
            time.sleep(0.3)
    return []


def consultar_recientes(nombre_tabla, limite=8):
    try:
        resultado = (
            supabase
            .table(nombre_tabla)
            .select("*")
            .order("id", desc=True)
            .limit(limite)
            .execute()
        )
        return resultado.data or []
    except Exception as error:
        print(f"Error consultando recientes de {nombre_tabla}: {error}")
        return []


def obtener_registro(nombre_tabla, registro_id):
    resultado = (
        supabase
        .table(nombre_tabla)
        .select("*")
        .eq("id", registro_id)
        .limit(1)
        .execute()
    )
    return resultado.data[0] if resultado.data else None


def estado_operativo(activo):
    disponibilidad = str(
        activo.get("disponibilidad") or ""
    ).strip().lower()
    estado = str(activo.get("estado") or "").strip().lower()
    return disponibilidad or estado


def registrar_movimiento(
    activo_id,
    persona_id,
    accion,
    observacion=None,
):
    return (
        supabase
        .table("movimientos")
        .insert({
            "activo_id": activo_id,
            "persona_id": persona_id,
            "accion": accion,
            "observacion": observacion or None,
        })
        .execute()
    )


def registrar_historial(activo_id, accion, detalle):
    """Guarda la accion en el historial junto con el usuario que la hizo."""
    registro = {
        "activo_id": activo_id,
        "accion": accion,
        "detalle": detalle,
    }

    if has_request_context() and current_user.is_authenticated:
        registro["usuario_id"] = int(current_user.id)
        registro["usuario_nombre"] = current_user.nombre

    try:
        return supabase.table("historial").insert(registro).execute()
    except Exception as error:
        # Si aun no existen las columnas usuario_*, guarda igual sin el usuario
        if "usuario_" in str(error):
            registro.pop("usuario_id", None)
            registro.pop("usuario_nombre", None)
            return supabase.table("historial").insert(registro).execute()
        raise


def _fecha_local(valor):
    """ISO en UTC -> '02/10 15:30' en hora Colombia."""
    try:
        dt = datetime.fromisoformat(str(valor).replace("Z", "+00:00"))
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return (dt - timedelta(hours=5)).strftime("%d/%m %H:%M")
    except (TypeError, ValueError):
        return str(valor or "")[:16].replace("T", " ")


def obtener_roles():
    return consultar_tabla("roles", "*", "nombre")


def plantilla_existe(nombre):
    return (Path(app.template_folder or "templates") / nombre).exists()


@app.context_processor
def utilidades_permisos():
    def tiene_permiso(modulo, accion="ver"):
        return (
            current_user.is_authenticated
            and current_user.tiene_permiso(modulo, accion)
        )

    return {
        "tiene_permiso": tiene_permiso,
    }


# ==================================================
# AUTENTICACION
# ==================================================

@app.route("/login", methods=["GET", "POST"])
def login():
    if current_user.is_authenticated:
        return redirect(url_for("dashboard"))

    if request.method == "POST":
        username = request.form.get("username", "").strip().lower()
        password = request.form.get("password", "")
        recordar = request.form.get("recordar") == "on"

        if not username or not password:
            flash("Escribe el usuario y la contrasena.", "danger")
            return render_template("login.html")

        try:
            usuario = autenticar_usuario(username, password)
            if usuario is None:
                flash("Usuario o contrasena incorrectos.", "danger")
                return render_template("login.html")

            login_user(usuario, remember=recordar)
            session.permanent = True
            session["ultimo_acceso"] = datetime.now(
                timezone.utc
            ).isoformat()
            flash(f"Bienvenido, {usuario.nombre}.", "success")

            siguiente = request.args.get("next", "")
            if siguiente.startswith("/") and not siguiente.startswith("//"):
                return redirect(siguiente)

            return redirect(url_for("dashboard"))

        except Exception as error:
            print(f"Error iniciando sesion: {error}")
            flash("No fue posible iniciar sesion.", "danger")

    return render_template("login.html")


@app.route("/logout")
@login_required
def logout():
    nombre = current_user.nombre
    return cerrar_sesion_y_redirigir(
        f"La sesion de {nombre} fue cerrada.",
        "success",
    )


# ==================================================
# DASHBOARD
# ==================================================

@app.route("/")
@login_required
@permiso_requerido("dashboard", "ver")
def dashboard():
    activos = consultar_tabla("activos", "*", "id")
    total = len(activos)
    disponibles = sum(
        1 for activo in activos
        if estado_operativo(activo) == "disponible"
    )
    asignados = sum(
        1 for activo in activos
        if estado_operativo(activo) == "asignado"
    )
    reparacion = sum(
        1 for activo in activos
        if "repar" in estado_operativo(activo)
    )
    bajas = sum(
        1 for activo in activos
        if "baja" in estado_operativo(activo)
    )

    # Activos por tipo (top 6)
    por_tipo = Counter(
        (str(a.get("tipo") or "").strip() or "Sin tipo")
        for a in activos
    ).most_common(6)

    # Activos asignados por area (top 5)
    por_area = Counter(
        str(a.get("area")).strip()
        for a in activos
        if estado_operativo(a) == "asignado" and a.get("area")
    ).most_common(5)

    # Tickets que siguen abiertos
    cerrados = {"cerrado", "resuelto", "solucionado", "finalizado", "completado"}
    tickets_abiertos = [
        t for t in consultar_tabla("tickets", "*", "id")
        if str(t.get("estado") or "").strip().lower() not in cerrados
    ]
    tickets_abiertos.sort(key=lambda t: t.get("id") or 0, reverse=True)

    codigos = {a.get("id"): a.get("codigo") for a in activos}

    # Actividad reciente: hora local y codigo del activo
    actividad = consultar_recientes("historial", 8)
    for mov in actividad:
        mov["fecha_local"] = _fecha_local(
            mov.get("created_at") or mov.get("fecha")
        )
        mov["activo_codigo"] = codigos.get(mov.get("activo_id"))

    lista_personas = consultar_tabla("personas", "*", "id")
    personas_activas = sum(
        1 for p in lista_personas
        if str(p.get("estado") or "Activo").strip().lower() == "activo"
    )

    return render_template(
        "dashboard.html",
        total=total,
        disponibles=disponibles,
        asignados=asignados,
        reparacion=reparacion,
        bajas=bajas,
        por_tipo=por_tipo,
        por_area=por_area,
        tickets_abiertos=tickets_abiertos[:5],
        total_tickets_abiertos=len(tickets_abiertos),
        codigos=codigos,
        personas_activas=personas_activas,
        total_actas=len(consultar_tabla("actas", "id")),
        actividad=actividad,
        hoy=datetime.now(timezone.utc) - timedelta(hours=5),
    )


# ==================================================
# INVENTARIO
# ==================================================

BUCKET_FACTURAS = "facturas"
EXT_FACTURA = {".pdf", ".jpg", ".jpeg", ".png"}
DISPONIBILIDADES = {
    "Disponible",
    "En reparación",
    "Inhabilitado",
    "Dado de baja",
}


def _es_laptop(tipo):
    t = (tipo or "").lower()
    return "laptop" in t or "portatil" in t or "portátil" in t


def subir_factura(archivo, carpeta):
    """Sube un archivo al bucket privado y devuelve la ruta (o None)."""
    if not archivo or not archivo.filename:
        return None

    ext = Path(archivo.filename).suffix.lower()
    if ext not in EXT_FACTURA:
        raise ValueError("La factura debe ser PDF, JPG o PNG.")

    ruta = f"{carpeta}/{uuid.uuid4().hex}{ext}"
    supabase.storage.from_(BUCKET_FACTURAS).upload(
        ruta,
        archivo.read(),
        {"content-type": archivo.mimetype or "application/octet-stream"},
    )
    return ruta


@app.route("/inventario/factura/<path:ruta>")
@login_required
@permiso_requerido("inventario", "ver")
def ver_factura(ruta):
    """Redirige a un enlace temporal (5 min) de la factura."""
    try:
        firmado = supabase.storage.from_(BUCKET_FACTURAS).create_signed_url(
            ruta, 300
        )
        url = firmado.get("signedURL") or firmado.get("signedUrl")
        if not url:
            raise ValueError("Sin URL")
        return redirect(url)
    except Exception as error:
        print(f"Error abriendo factura: {error}")
        flash("No fue posible abrir la factura.", "danger")
        return redirect(url_for("inventario"))


@app.route("/inventario")
@login_required
@permiso_requerido("inventario", "ver")
def inventario():
    activos = consultar_tabla("activos", "*", "codigo")

    personas_activas = [
        p for p in consultar_tabla("personas", "*", "nombre")
        if str(p.get("estado") or "Activo").strip().lower() == "activo"
    ]

    adicionales = {}
    for ad in consultar_tabla("activos_adicionales", "*", "id"):
        ad["factura_url"] = (
            url_for("ver_factura", ruta=ad["factura_ruta"])
            if ad.get("factura_ruta") else ""
        )
        adicionales.setdefault(ad["activo_id"], []).append(ad)

    return render_template(
        "activos.html",
        activos=activos,
        personas=personas_activas,
        adicionales=adicionales,
    )


def _campos_activo():
    """Lee y limpia los campos del formulario de crear/editar activo."""
    f = request.form

    def texto(nombre):
        return f.get(nombre, "").strip() or None

    return {
        "codigo": f.get("codigo", "").strip(),
        "tipo": f.get("tipo", "").strip(),
        "marca": texto("marca"),
        "modelo": texto("modelo"),
        "serial": texto("serial"),
        "estado": texto("estado"),  # estado fisico
        "cantidad": f.get("cantidad", type=int),
        "procesador": texto("procesador"),
        "ram": texto("ram"),
        "disco": texto("disco"),
        "hostname": texto("hostname"),
        "ip": texto("ip"),
        "mac": texto("mac"),
        "sistema_operativo": texto("sistema_operativo"),
        "fecha_compra": texto("fecha_compra"),
        "garantia": texto("garantia"),
        "observaciones": texto("observaciones"),
    }


def _codigo_en_uso(codigo, excluir_id=None):
    consulta = (
        supabase
        .table("activos")
        .select("id")
        .eq("codigo", codigo)
    )
    if excluir_id is not None:
        consulta = consulta.neq("id", excluir_id)
    return bool(consulta.limit(1).execute().data)


def _persona_para_asignar():
    """Devuelve la persona elegida en el formulario (o None)."""
    persona_id = request.form.get("persona_id", type=int)
    if not persona_id:
        return None

    persona = obtener_registro("personas", persona_id)
    if persona is None:
        raise ValueError("La persona seleccionada no existe.")
    if str(persona.get("estado") or "Activo").strip().lower() != "activo":
        raise ValueError("La persona seleccionada está inactiva.")
    return persona


def _asignar_a_persona(activo_id, codigo, persona, anterior=None):
    """Asigna el activo (nombre y area salen de la persona)."""
    supabase.table("activos").update({
        "disponibilidad": "Asignado",
        "asignado_a": persona.get("nombre"),
        "area": persona.get("area"),
    }).eq("id", activo_id).execute()

    registrar_movimiento(activo_id, persona["id"], "Asignacion", None)

    detalle = f"Activo {codigo} asignado a {persona.get('nombre')}."
    if anterior:
        detalle += f" Antes: {anterior}."
    registrar_historial(activo_id, "Asignacion", detalle)


def _guardar_archivos(activo_id, tipo):
    """Sube la factura y los adicionales. Devuelve una lista de avisos."""
    avisos = []

    try:
        ruta = subir_factura(
            request.files.get("factura"),
            f"activos/{activo_id}",
        )
        if ruta:
            supabase.table("activos").update(
                {"factura_ruta": ruta}
            ).eq("id", activo_id).execute()
    except Exception as error:
        print(f"Error subiendo factura: {error}")
        avisos.append(f"No se pudo guardar la factura: {error}")

    if _es_laptop(tipo):
        tipos = request.form.getlist("adicional_tipo")
        descs = request.form.getlist("adicional_descripcion")
        archivos = request.files.getlist("adicional_factura")

        for i, desc in enumerate(descs):
            desc = desc.strip()
            if not desc:
                continue
            try:
                ruta = subir_factura(
                    archivos[i] if i < len(archivos) else None,
                    f"activos/{activo_id}/adicionales",
                )
                supabase.table("activos_adicionales").insert({
                    "activo_id": activo_id,
                    "tipo": tipos[i] if i < len(tipos) else None,
                    "descripcion": desc,
                    "factura_ruta": ruta,
                }).execute()
            except Exception as error:
                print(f"Error guardando adicional: {error}")
                avisos.append(
                    f"No se pudo guardar el adicional '{desc}': {error}"
                )

    return avisos


@app.route("/inventario/crear", methods=["POST"])
@login_required
@permiso_requerido("inventario", "crear")
def crear_activo_web():
    datos = _campos_activo()

    if not datos["codigo"] or not datos["tipo"]:
        flash("El codigo y el tipo son obligatorios.", "danger")
        return redirect(url_for("inventario"))

    try:
        if _codigo_en_uso(datos["codigo"]):
            raise ValueError(
                f"Ya existe un activo con el codigo {datos['codigo']}."
            )

        persona = _persona_para_asignar()

        disp = request.form.get("disponibilidad", "Disponible")
        if disp not in DISPONIBILIDADES:
            disp = "Disponible"
        # Si se eligio persona, queda Asignado al terminar de crear
        datos["disponibilidad"] = "Disponible" if persona else disp

        nuevo = supabase.table("activos").insert(datos).execute()
        activo = nuevo.data[0]

        registrar_historial(
            activo["id"],
            "Creacion",
            f"Activo {datos['codigo']} registrado en el inventario.",
        )

        if persona:
            _asignar_a_persona(activo["id"], datos["codigo"], persona)

        for aviso in _guardar_archivos(activo["id"], datos["tipo"]):
            flash(aviso, "warning")

        flash(f"Activo {datos['codigo']} creado correctamente.", "success")
    except Exception as error:
        print(f"Error creando activo: {error}")
        flash(f"No fue posible crear el activo: {error}", "danger")

    return redirect(url_for("inventario"))


@app.route("/inventario/<int:activo_id>/editar", methods=["POST"])
@login_required
@permiso_requerido("inventario", "editar")
def editar_activo_web(activo_id):
    datos = _campos_activo()

    if not datos["codigo"] or not datos["tipo"]:
        flash("El codigo y el tipo son obligatorios.", "danger")
        return redirect(url_for("inventario"))

    try:
        anterior = obtener_registro("activos", activo_id)
        if anterior is None:
            raise ValueError("El activo no existe.")

        if _codigo_en_uso(datos["codigo"], excluir_id=activo_id):
            raise ValueError(
                f"Otro activo ya usa el codigo {datos['codigo']}."
            )

        persona = _persona_para_asignar()
        esta_asignado = estado_operativo(anterior) == "asignado"

        # La disponibilidad solo se edita a mano si el activo no esta asignado
        if not esta_asignado and not persona:
            disp = request.form.get("disponibilidad", "")
            if disp in DISPONIBILIDADES:
                datos["disponibilidad"] = disp

        supabase.table("activos").update(datos).eq("id", activo_id).execute()

        detalle = f"Datos del activo {datos['codigo']} actualizados."
        if anterior.get("codigo") != datos["codigo"]:
            detalle = (
                f"Datos actualizados. Codigo: "
                f"{anterior.get('codigo')} -> {datos['codigo']}."
            )
        registrar_historial(activo_id, "Edicion", detalle)

        # Asignar o reasignar solo si cambio la persona.
        # Para quitar la asignacion se usa Devoluciones.
        if persona and persona.get("nombre") != anterior.get("asignado_a"):
            _asignar_a_persona(
                activo_id,
                datos["codigo"],
                persona,
                anterior=anterior.get("asignado_a") if esta_asignado else None,
            )

        for aviso in _guardar_archivos(activo_id, datos["tipo"]):
            flash(aviso, "warning")

        flash(f"Activo {datos['codigo']} actualizado.", "success")
    except Exception as error:
        print(f"Error editando activo: {error}")
        flash(f"No fue posible actualizar el activo: {error}", "danger")

    return redirect(url_for("inventario"))


def _cambiar_estado_activo(
    activo_id,
    estado_nuevo,
    accion,
    motivo_obligatorio=False,
):
    """Inhabilita o da de baja un activo que no este asignado."""
    motivo = request.form.get("motivo", "").strip()

    if motivo_obligatorio and not motivo:
        flash("El motivo es obligatorio.", "danger")
        return redirect(url_for("inventario"))

    try:
        activo = obtener_registro("activos", activo_id)
        if activo is None:
            raise ValueError("El activo no existe.")

        actual = estado_operativo(activo)
        if actual == "asignado":
            raise ValueError(
                "El activo esta asignado. Registra primero la devolucion."
            )
        if "baja" in actual or "inhabilit" in actual:
            raise ValueError("El activo ya esta fuera del inventario operativo.")

        supabase.table("activos").update({
            "disponibilidad": estado_nuevo,
        }).eq("id", activo_id).execute()

        codigo = activo.get("codigo")
        detalle = f"Activo {codigo}: {accion.lower()}."
        if motivo:
            detalle += f" Motivo: {motivo}"
        registrar_historial(activo_id, accion, detalle)

        flash(f"{codigo}: {accion.lower()} registrada.", "success")
    except Exception as error:
        print(f"Error en {accion.lower()} de activo: {error}")
        flash(str(error), "danger")

    return redirect(url_for("inventario"))


@app.route("/inventario/<int:activo_id>/inhabilitar", methods=["POST"])
@login_required
@permiso_requerido("inventario", "editar")
def inhabilitar_activo_web(activo_id):
    return _cambiar_estado_activo(
        activo_id,
        "Inhabilitado",
        "Inhabilitacion",
    )


@app.route("/inventario/<int:activo_id>/baja", methods=["POST"])
@login_required
@permiso_requerido("inventario", "eliminar")
def dar_de_baja_activo_web(activo_id):
    return _cambiar_estado_activo(
        activo_id,
        "Dado de baja",
        "Baja",
        motivo_obligatorio=True,
    )


@app.route("/inventario/<int:activo_id>/reactivar", methods=["POST"])
@login_required
@permiso_requerido("inventario", "eliminar")  # antes: "editar"
def reactivar_activo_web(activo_id):
    try:
        activo = obtener_registro("activos", activo_id)
        if activo is None:
            raise ValueError("El activo no existe.")

        actual = estado_operativo(activo)
        if "baja" not in actual and "inhabilit" not in actual:
            raise ValueError("Solo se pueden reactivar activos inhabilitados o dados de baja.")

        supabase.table("activos").update({
            "disponibilidad": "Disponible",
        }).eq("id", activo_id).execute()

        codigo = activo.get("codigo")
        registrar_historial(
            activo_id,
            "Reactivacion",
            f"Activo {codigo} reactivado. Estado anterior: "
            f"{activo.get('disponibilidad') or activo.get('estado')}.",
        )
        flash(f"{codigo} reactivado y disponible.", "success")
    except Exception as error:
        print(f"Error reactivando activo: {error}")
        flash(str(error), "danger")

    return redirect(url_for("inventario"))


# ==================================================
# PERSONAS
# ==================================================

@app.route("/personas")
@login_required
@permiso_requerido("personas", "ver")
def personas():
    lista_personas = consultar_tabla("personas", "*", "nombre")
    return render_template("personas.html", personas=lista_personas)


@app.route("/persona/crear", methods=["POST"])
@app.route("/personas/crear", methods=["POST"])
@login_required
@permiso_requerido("personas", "crear")
def crear_persona_web():
    nombre = request.form.get("nombre", "").strip()
    documento = request.form.get("documento", "").strip()
    correo = request.form.get("correo", "").strip()
    cargo = request.form.get("cargo", "").strip()
    area = request.form.get("area", "").strip()

    if not nombre:
        flash("El nombre es obligatorio.", "danger")
        return redirect(url_for("personas"))

    try:
        supabase.table("personas").insert({
            "nombre": nombre,
            "documento": documento or None,
            "correo": correo or None,
            "cargo": cargo or None,
            "area": area or None,
            "estado": "Activo",
        }).execute()
        flash(f"Persona {nombre} creada correctamente.", "success")
    except Exception as error:
        print(f"Error creando persona: {error}")
        flash(f"No fue posible crear la persona: {error}", "danger")

    return redirect(url_for("personas"))


@app.route("/personas/<int:persona_id>/editar", methods=["POST"])
@login_required
@permiso_requerido("personas", "editar")
def editar_persona_web(persona_id):
    nombre = request.form.get("nombre", "").strip()
    documento = request.form.get("documento", "").strip()
    correo = request.form.get("correo", "").strip()
    cargo = request.form.get("cargo", "").strip()
    area = request.form.get("area", "").strip()
    estado = request.form.get("estado", "Activo").strip()

    if not nombre:
        flash("El nombre es obligatorio.", "danger")
        return redirect(url_for("persona_detalle", persona_id=persona_id))

    if estado not in {"Activo", "Inactivo"}:
        estado = "Activo"

    try:
        anterior = obtener_registro("personas", persona_id)
        supabase.table("personas").update({
            "nombre": nombre,
            "documento": documento or None,
            "correo": correo or None,
            "cargo": cargo or None,
            "area": area or None,
            "estado": estado,
        }).eq("id", persona_id).execute()

        # Los activos se vinculan por nombre: si cambia, se mantiene el vínculo
        if anterior and anterior.get("nombre") and anterior["nombre"] != nombre:
            supabase.table("activos").update({
                "asignado_a": nombre,
            }).eq("asignado_a", anterior["nombre"]).execute()

        flash("Persona actualizada correctamente.", "success")
    except Exception as error:
        print(f"Error actualizando persona: {error}")
        flash(f"No fue posible actualizar la persona: {error}", "danger")

    return redirect(url_for("persona_detalle", persona_id=persona_id))


@app.route("/personas/<int:persona_id>/desactivar", methods=["POST"])
@login_required
@permiso_requerido("personas", "eliminar")
def desactivar_persona_web(persona_id):
    try:
        supabase.table("personas").update({
            "estado": "Inactivo"
        }).eq("id", persona_id).execute()
        flash("Persona desactivada correctamente.", "success")
    except Exception as error:
        print(f"Error desactivando persona: {error}")
        flash(f"No fue posible desactivar la persona: {error}", "danger")

    destino = request.form.get("volver")
    if destino == "detalle":
        return redirect(url_for("persona_detalle", persona_id=persona_id))
    return redirect(url_for("personas"))

@app.route("/personas/<int:persona_id>/activar", methods=["POST"])
@login_required
@permiso_requerido("personas", "editar")
def activar_persona_web(persona_id):
    try:
        supabase.table("personas").update({
            "estado": "Activo"
        }).eq("id", persona_id).execute()
        flash("Persona activada correctamente.", "success")
    except Exception as error:
        print(f"Error activando persona: {error}")
        flash(f"No fue posible activar la persona: {error}", "danger")

    destino = request.form.get("volver")
    if destino == "detalle":
        return redirect(url_for("persona_detalle", persona_id=persona_id))
    return redirect(url_for("personas"))

@app.route("/personas/<int:persona_id>")
@login_required
@permiso_requerido("personas", "ver")
def persona_detalle(persona_id):
    persona = obtener_registro("personas", persona_id)
    if persona is None:
        abort(404)

    nombre = str(persona.get("nombre") or "").strip()

    # Activos asignados actualmente
    try:
        activos_asignados = (
            supabase.table("activos")
            .select("*")
            .eq("asignado_a", nombre)
            .order("codigo")
            .execute()
        ).data or []
    except Exception as error:
        print(f"Error consultando activos de persona: {error}")
        activos_asignados = []

    # Historial de movimientos (asignaciones / devoluciones)
    try:
        movimientos = (
            supabase.table("movimientos")
            .select("*")
            .eq("persona_id", persona_id)
            .order("id", desc=True)
            .limit(50)
            .execute()
        ).data or []
    except Exception as error:
        print(f"Error consultando movimientos: {error}")
        movimientos = []

    # Mapa id -> código de activo para mostrar el código en los movimientos
    codigos = {}
    ids = {m.get("activo_id") for m in movimientos if m.get("activo_id")}
    if ids:
        try:
            filas = (
                supabase.table("activos")
                .select("id,codigo")
                .in_("id", list(ids))
                .execute()
            ).data or []
            codigos = {f["id"]: f["codigo"] for f in filas}
        except Exception as error:
            print(f"Error consultando códigos: {error}")

    # Actas de la persona
    try:
        actas_persona = (
            supabase.table("actas")
            .select("*")
            .eq("persona_id", persona_id)
            .order("id", desc=True)
            .execute()
        ).data or []
    except Exception as error:
        print(f"Error consultando actas: {error}")
        actas_persona = []

    return render_template(
        "persona_detalle.html",
        persona=persona,
        activos=activos_asignados,
        movimientos=movimientos,
        codigos=codigos,
        actas=actas_persona,
    )


# ==================================================
# ASIGNACIONES
# ==================================================

@app.route("/asignaciones")
@login_required
@permiso_requerido("asignaciones", "ver")
def asignaciones():
    activos = consultar_tabla("activos", "*", "codigo")
    lista_personas = consultar_tabla("personas", "*", "nombre")

    activos_disponibles = [
        activo for activo in activos
        if estado_operativo(activo) == "disponible"
    ]
    personas_activas = [
        persona for persona in lista_personas
        if str(persona.get("estado") or "Activo").strip().lower()
        == "activo"
    ]

    return render_template(
        "asignaciones.html",
        activos=activos_disponibles,
        personas=personas_activas,
    )


@app.route("/asignar", methods=["POST"])
@app.route("/asignaciones/crear", methods=["POST"])
@login_required
@permiso_requerido("asignaciones", "crear")
def crear_asignacion_web():
    activo_id = request.form.get("activo_id", type=int)
    persona_id = request.form.get("persona_id", type=int)
    observacion = request.form.get("observacion", "").strip()

    if not activo_id or not persona_id:
        flash("Selecciona una persona y un activo.", "danger")
        return redirect(url_for("asignaciones"))

    try:
        activo = obtener_registro("activos", activo_id)
        persona = obtener_registro("personas", persona_id)

        if activo is None:
            raise ValueError("El activo seleccionado no existe.")
        if persona is None:
            raise ValueError("La persona seleccionada no existe.")
        if estado_operativo(activo) != "disponible":
            raise ValueError("El activo ya no esta disponible.")

        nombre_persona = persona.get("nombre")
        area_persona = persona.get("area")
        codigo_activo = activo.get("codigo")

        supabase.table("activos").update({
            "disponibilidad": "Asignado",
            "asignado_a": nombre_persona,
            "area": area_persona,
            "observaciones": observacion or None,
        }).eq("id", activo_id).execute()

        registrar_movimiento(
            activo_id,
            persona_id,
            "Asignacion",
            observacion,
        )
        registrar_historial(
            activo_id,
            "Asignacion",
            f"Activo {codigo_activo} asignado a {nombre_persona}.",
        )
        flash(
            f"{codigo_activo} asignado correctamente a {nombre_persona}.",
            "success",
        )
    except Exception as error:
        print(f"Error asignando activo: {error}")
        flash(str(error), "danger")

    return redirect(url_for("asignaciones"))


# ==================================================
# DEVOLUCIONES
# ==================================================

@app.route("/devoluciones")
@login_required
@permiso_requerido("devoluciones", "ver")
def devoluciones():
    activos = consultar_tabla("activos", "*", "codigo")
    activos_asignados = [
        activo for activo in activos
        if estado_operativo(activo) == "asignado"
    ]
    return render_template("devoluciones.html", activos=activos_asignados)


@app.route("/devolver/<int:activo_id>", methods=["POST"])
@app.route("/devoluciones/<int:activo_id>", methods=["POST"])
@login_required
@permiso_requerido("devoluciones", "crear")
def devolver_activo_web(activo_id):
    estado_destino = request.form.get("estado_destino", "Disponible")
    observacion = request.form.get("observacion", "").strip()

    estados_permitidos = {
        "Disponible",
        "En reparacion",
        "En reparación",
        "Dado de baja",
    }

    if estado_destino not in estados_permitidos:
        flash("El estado de devolucion no es valido.", "danger")
        return redirect(url_for("devoluciones"))

    try:
        activo = obtener_registro("activos", activo_id)
        if activo is None:
            raise ValueError("El activo seleccionado no existe.")
        if estado_operativo(activo) != "asignado":
            raise ValueError("El activo seleccionado no esta asignado.")

        codigo_activo = activo.get("codigo")
        persona_anterior = activo.get("asignado_a")

        supabase.table("activos").update({
            "disponibilidad": estado_destino,
            "asignado_a": None,
            "area": None,
            "observaciones": observacion or None,
        }).eq("id", activo_id).execute()

        registrar_movimiento(
            activo_id,
            None,
            "Devolucion",
            observacion,
        )
        registrar_historial(
            activo_id,
            "Devolucion",
            (
                f"Activo {codigo_activo} devuelto por "
                f"{persona_anterior or 'persona no identificada'}. "
                f"Nuevo estado: {estado_destino}."
            ),
        )
        flash(f"{codigo_activo} fue devuelto correctamente.", "success")
    except Exception as error:
        print(f"Error devolviendo activo: {error}")
        flash(str(error), "danger")

    return redirect(url_for("devoluciones"))


# ==================================================
# TICKETS
# ==================================================

@app.route("/tickets")
@login_required
@permiso_requerido("tickets", "ver")
def tickets():
    lista_tickets = consultar_tabla("tickets", "*", "id")
    lista_tickets.sort(key=lambda t: t.get("id") or 0, reverse=True)
    activos = consultar_tabla("activos", "*", "codigo")
    return render_template(
        "tickets.html",
        tickets=lista_tickets,
        activos=activos,
    )


@app.route("/ticket/crear", methods=["POST"])
@app.route("/tickets/crear", methods=["POST"])
@login_required
@permiso_requerido("tickets", "crear")
def crear_ticket_web():
    activo_id = request.form.get("activo_id", type=int)
    titulo = request.form.get("titulo", "").strip()
    descripcion = request.form.get("descripcion", "").strip()

    if not activo_id or not titulo or not descripcion:
        flash("El activo, el titulo y la descripcion son obligatorios.", "danger")
        return redirect(url_for("tickets"))

    try:
        supabase.table("tickets").insert({
            "activo_id": activo_id,
            "titulo": titulo,
            "descripcion": descripcion,
            "estado": "Pendiente",
        }).execute()
        flash("Ticket creado correctamente.", "success")
    except Exception as error:
        print(f"Error creando ticket: {error}")
        flash(f"No fue posible crear el ticket: {error}", "danger")

    return redirect(url_for("tickets"))


ESTADOS_TICKET = {
    "Pendiente",
    "En proceso",
    "Esperando repuesto",
    "Solucionado",
}


@app.route("/tickets/<int:ticket_id>/estado", methods=["POST"])
@login_required
@permiso_requerido("tickets", "editar")
def cambiar_estado_ticket_web(ticket_id):
    estado = request.form.get("estado", "").strip()

    if estado not in ESTADOS_TICKET:
        flash("El estado del ticket no es valido.", "danger")
        return redirect(url_for("tickets"))

    try:
        if obtener_registro("tickets", ticket_id) is None:
            raise ValueError("El ticket no existe.")

        supabase.table("tickets").update({
            "estado": estado,
        }).eq("id", ticket_id).execute()

        # Opcional: guarda la fecha de solucion si la columna existe
        # (ver SQL en las notas). Si no existe, se ignora sin afectar.
        try:
            supabase.table("tickets").update({
                "fecha_solucion": (
                    datetime.now(timezone.utc).isoformat()
                    if estado == "Solucionado" else None
                ),
            }).eq("id", ticket_id).execute()
        except Exception:
            pass

        if estado == "Solucionado":
            flash(f"Ticket #{ticket_id} marcado como solucionado.", "success")
        else:
            flash(f"Ticket #{ticket_id} actualizado a {estado}.", "success")
    except Exception as error:
        print(f"Error actualizando ticket: {error}")
        flash(str(error), "danger")

    return redirect(url_for("tickets"))


# ==================================================
# ACTAS PDF
# ==================================================

@app.route("/actas")
@login_required
@permiso_requerido("actas", "ver")
def actas():
    lista_actas = consultar_tabla("actas", "*", "id")
    lista_personas = consultar_tabla("personas", "*", "nombre")
    return render_template(
        "actas.html",
        actas=lista_actas,
        personas=lista_personas,
    )


@app.route("/actas/generar", methods=["POST"])
@login_required
@permiso_requerido("actas", "crear")
def generar_acta_web():
    persona_id = request.form.get("persona_id", type=int)
    tipo = request.form.get("tipo", "Entrega").strip()
    observaciones = request.form.get("observaciones", "").strip()

    if not persona_id:
        flash("Debes seleccionar una persona.", "danger")
        return redirect(url_for("actas"))

    tipos_permitidos = {
        "Entrega",
        "Devolucion",
        "Cambio",
        "Paz y salvo",
    }
    if tipo not in tipos_permitidos:
        flash("El tipo de acta no es valido.", "danger")
        return redirect(url_for("actas"))

    try:
        persona = obtener_registro("personas", persona_id)
        if persona is None:
            raise ValueError("La persona seleccionada no existe.")

        nombre_persona = str(persona.get("nombre") or "").strip()
        if not nombre_persona:
            raise ValueError("La persona no tiene un nombre valido.")

        respuesta_activos = (
            supabase
            .table("activos")
            .select("*")
            .eq("asignado_a", nombre_persona)
            .order("codigo")
            .execute()
        )
        activos_persona = respuesta_activos.data or []

        ruta_pdf, archivo_pdf = generar_acta_pdf(
            persona=persona,
            activos=activos_persona,
            tipo=tipo,
            observaciones=observaciones,
        )

        supabase.table("actas").insert({
            "persona_id": persona_id,
            "tipo": tipo,
            "ruta_pdf": archivo_pdf,
            "observaciones": observaciones or None,
        }).execute()

        print(f"Acta guardada en: {ruta_pdf}")
        return redirect(
            url_for(
                "descargar_acta",
                nombre_archivo=archivo_pdf,
            )
        )
    except Exception as error:
        print(f"Error generando acta PDF: {error}")
        flash(f"No fue posible generar el acta: {error}", "danger")
        return redirect(url_for("actas"))


@app.route("/actas/descargar/<path:nombre_archivo>")
@login_required
@permiso_requerido("actas", "ver")
def descargar_acta(nombre_archivo):
    nombre_limpio = Path(nombre_archivo).name
    carpeta_pdf = Path("pdf").resolve()
    archivo_pdf = carpeta_pdf / nombre_limpio

    if not archivo_pdf.exists():
        flash("El archivo PDF solicitado no existe.", "danger")
        return redirect(url_for("actas"))

    return send_from_directory(
        str(carpeta_pdf),
        nombre_limpio,
        as_attachment=True,
    )


# ==================================================
# REPORTES Y CONFIGURACION
# ==================================================

def _sello_archivo():
    """Fecha de hoy (hora Colombia) para nombrar los archivos."""
    return (datetime.now(timezone.utc) - timedelta(hours=5)).strftime("%Y%m%d")


def respuesta_excel(nombre_archivo, hoja, encabezados, filas):
    """Construye un .xlsx en memoria y lo devuelve como descarga."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    libro = Workbook()
    ws = libro.active
    ws.title = hoja[:31]

    ws.append(list(encabezados))
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F3A5F")
        celda.alignment = Alignment(vertical="center")

    anchos = [len(str(h)) for h in encabezados]
    for fila in filas:
        valores = ["" if v is None else v for v in fila]
        ws.append(valores)
        for i, valor in enumerate(valores):
            anchos[i] = max(anchos[i], len(str(valor)))

    # Un texto que empieza con "=" no debe interpretarse como formula
    for fila_celdas in ws.iter_rows(min_row=2):
        for celda in fila_celdas:
            if isinstance(celda.value, str) and celda.value.startswith("="):
                celda.data_type = "s"

    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(ancho, 50) + 2

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    buffer = io.BytesIO()
    libro.save(buffer)
    buffer.seek(0)

    return send_file(
        buffer,
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        as_attachment=True,
        download_name=nombre_archivo,
    )


def _texto_fecha(valor):
    """'2026-09-30T20:53:27+00:00' -> '2026-09-30 20:53'."""
    return str(valor or "")[:16].replace("T", " ")


@app.route("/reportes")
@login_required
@permiso_requerido("reportes", "ver")
def reportes():
    resumen = {
        "activos": len(consultar_tabla("activos", "id")),
        "personas": len(consultar_tabla("personas", "id")),
        "tickets": len(consultar_tabla("tickets", "id")),
        "actas": len(consultar_tabla("actas", "id")),
        "historial": len(consultar_tabla("historial", "id")),
    }
    return render_template("reportes.html", resumen=resumen)


@app.route("/reportes/inventario.xlsx")
@login_required
@permiso_requerido("reportes", "ver")
def reporte_inventario():
    filtro = request.args.get("estado", "todos").strip().lower()
    activos = consultar_tabla("activos", "*", "codigo")

    if filtro in {"disponible", "asignado", "repar", "baja"}:
        activos = [a for a in activos if filtro in estado_operativo(a)]

    filas = [
        [
            a.get("codigo"),
            a.get("tipo"),
            a.get("marca"),
            a.get("modelo"),
            a.get("serial"),
            a.get("disponibilidad") or a.get("estado"),
            a.get("asignado_a"),
            a.get("area"),
            a.get("observaciones"),
        ]
        for a in activos
    ]

    try:
        return respuesta_excel(
            f"inventario_{_sello_archivo()}.xlsx",
            "Inventario",
            ["Codigo", "Tipo", "Marca", "Modelo", "Serial",
             "Estado", "Asignado a", "Area", "Observaciones"],
            filas,
        )
    except ImportError:
        flash("Falta instalar openpyxl (pip install openpyxl).", "danger")
        return redirect(url_for("reportes"))


@app.route("/reportes/personas.xlsx")
@login_required
@permiso_requerido("reportes", "ver")
def reporte_personas():
    lista = consultar_tabla("personas", "*", "nombre")
    filas = [
        [
            p.get("nombre"), p.get("documento"), p.get("correo"),
            p.get("cargo"), p.get("area"), p.get("estado") or "Activo",
        ]
        for p in lista
    ]

    try:
        return respuesta_excel(
            f"personas_{_sello_archivo()}.xlsx",
            "Personas",
            ["Nombre", "Documento", "Correo", "Cargo", "Area", "Estado"],
            filas,
        )
    except ImportError:
        flash("Falta instalar openpyxl (pip install openpyxl).", "danger")
        return redirect(url_for("reportes"))


@app.route("/reportes/tickets.xlsx")
@login_required
@permiso_requerido("reportes", "ver")
def reporte_tickets():
    filtro = request.args.get("estado", "todos").strip().lower()
    lista = consultar_tabla("tickets", "*", "id")
    codigos = {
        a.get("id"): a.get("codigo")
        for a in consultar_tabla("activos", "id,codigo")
    }

    def resuelto(t):
        e = str(t.get("estado") or "").lower()
        return "solucion" in e or "cerrado" in e or "resuelto" in e

    if filtro == "abiertos":
        lista = [t for t in lista if not resuelto(t)]
    elif filtro == "solucionados":
        lista = [t for t in lista if resuelto(t)]

    filas = [
        [
            t.get("id"),
            codigos.get(t.get("activo_id")) or t.get("activo_id"),
            t.get("titulo"),
            t.get("descripcion"),
            t.get("estado"),
            _texto_fecha(t.get("created_at") or t.get("fecha")),
            _texto_fecha(t.get("fecha_solucion")),
        ]
        for t in sorted(lista, key=lambda t: t.get("id") or 0, reverse=True)
    ]

    try:
        return respuesta_excel(
            f"tickets_{_sello_archivo()}.xlsx",
            "Tickets",
            ["ID", "Activo", "Titulo", "Descripcion", "Estado",
             "Fecha", "Fecha solucion"],
            filas,
        )
    except ImportError:
        flash("Falta instalar openpyxl (pip install openpyxl).", "danger")
        return redirect(url_for("reportes"))


@app.route("/reportes/historial.xlsx")
@login_required
@permiso_requerido("reportes", "ver")
def reporte_historial():
    desde = request.args.get("desde", "").strip()
    hasta = request.args.get("hasta", "").strip()

    lista = consultar_tabla("historial", "*", "id")
    codigos = {
        a.get("id"): a.get("codigo")
        for a in consultar_tabla("activos", "id,codigo")
    }

    def dentro_del_rango(fila):
        dia = str(fila.get("created_at") or fila.get("fecha") or "")[:10]
        if not dia:
            return True
        if desde and dia < desde:
            return False
        if hasta and dia > hasta:
            return False
        return True

    if desde or hasta:
        lista = [h for h in lista if dentro_del_rango(h)]

    filas = [
        [
            _texto_fecha(h.get("created_at") or h.get("fecha")),
            codigos.get(h.get("activo_id")) or h.get("activo_id"),
            h.get("accion"),
            h.get("detalle"),
        ]
        for h in sorted(lista, key=lambda h: h.get("id") or 0, reverse=True)
    ]

    try:
        return respuesta_excel(
            f"historial_{_sello_archivo()}.xlsx",
            "Historial",
            ["Fecha", "Activo", "Accion", "Detalle"],
            filas,
        )
    except ImportError:
        flash("Falta instalar openpyxl (pip install openpyxl).", "danger")
        return redirect(url_for("reportes"))


@app.route("/reportes/actas.zip")
@login_required
@permiso_requerido("reportes", "ver")
def reporte_actas():
    carpeta = Path("pdf").resolve()
    buffer = io.BytesIO()
    incluidos = set()

    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for acta in consultar_tabla("actas", "*", "id"):
            nombre = Path(str(acta.get("ruta_pdf") or "")).name
            if not nombre or nombre in incluidos:
                continue
            ruta = carpeta / nombre
            if ruta.exists():
                zf.write(ruta, nombre)
                incluidos.add(nombre)

    if not incluidos:
        flash("No hay archivos PDF de actas para descargar.", "warning")
        return redirect(url_for("reportes"))

    buffer.seek(0)
    return send_file(
        buffer,
        mimetype="application/zip",
        as_attachment=True,
        download_name=f"actas_{_sello_archivo()}.zip",
    )


# --------------------------------------------------
# CONFIGURACION
# --------------------------------------------------

CONFIG_DEFECTO = {
    "nombre_empresa": "Editorial Planeta",
    "correo_soporte": "",
    "dias_alerta_garantia": 30,
}


def obtener_config():
    try:
        r = (
            supabase.table("configuracion")
            .select("*")
            .eq("id", 1)
            .limit(1)
            .execute()
        )
        if r.data:
            return {**CONFIG_DEFECTO, **r.data[0]}
    except Exception as error:
        print(f"Error leyendo configuracion: {error}")
    return dict(CONFIG_DEFECTO)


@app.route("/configuracion", methods=["GET", "POST"])
@login_required
@permiso_requerido("configuracion", "ver")
def configuracion():
    cfg = obtener_config()
    errores = {}

    if request.method == "POST":
        if not current_user.tiene_permiso("configuracion", "editar"):
            abort(403)

        nombre = request.form.get("nombre_empresa", "").strip()
        correo = request.form.get("correo_soporte", "").strip().lower()
        dias_txt = request.form.get("dias_alerta_garantia", "").strip()

        if not nombre:
            errores["nombre_empresa"] = "El nombre es obligatorio."
        if correo and not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", correo):
            errores["correo_soporte"] = "Correo inválido."
        try:
            dias = int(dias_txt)
            if not 1 <= dias <= 365:
                raise ValueError
        except ValueError:
            dias = dias_txt
            errores["dias_alerta_garantia"] = "Debe ser un número entre 1 y 365."

        cfg.update({
            "nombre_empresa": nombre,
            "correo_soporte": correo,
            "dias_alerta_garantia": dias,
        })

        if not errores:
            try:
                supabase.table("configuracion").upsert({
                    "id": 1,
                    "nombre_empresa": nombre,
                    "correo_soporte": correo or None,
                    "dias_alerta_garantia": dias,
                }).execute()
                flash("Configuración guardada correctamente.", "success")
                return redirect(url_for("configuracion"))
            except Exception as error:
                print(f"Error guardando configuracion: {error}")
                flash(f"No fue posible guardar: {error}", "danger")

    return render_template("configuracion.html", cfg=cfg, errores=errores)


# ==================================================
# ADMINISTRACION DE USUARIOS
# ==================================================

@app.route("/usuarios")
@login_required
@solo_administrador
def usuarios():
    lista_usuarios = consultar_tabla("usuarios", "*", "nombre")
    roles = consultar_tabla("roles", "*", "nombre")

    if plantilla_existe("usuarios/lista.html"):
        return render_template(
            "usuarios/lista.html",
            usuarios=lista_usuarios,
            roles=roles,
        )

    return {
        "usuarios": lista_usuarios,
        "roles": roles,
    }


@app.route("/usuarios/crear", methods=["POST"])
@login_required
@solo_administrador
def crear_usuario_web():
    nombre = request.form.get("nombre", "").strip()
    username = request.form.get("username", "").strip().lower()
    correo = request.form.get("correo", "").strip().lower()
    password = request.form.get("password", "")
    rol_id = request.form.get("rol_id", type=int)

    try:
        crear_usuario(
            nombre=nombre,
            username=username,
            correo=correo,
            password=password,
            rol_id=rol_id,
        )
        flash("Usuario creado correctamente.", "success")
    except Exception as error:
        print(f"Error creando usuario: {error}")
        flash(str(error), "danger")

    return redirect(url_for("usuarios"))


@app.route(
    "/usuarios/<int:usuario_id>/editar",
    methods=["POST"],
)
@login_required
@solo_administrador
def editar_usuario_web(usuario_id):
    nombre = request.form.get("nombre", "").strip()
    username = request.form.get("username", "").strip().lower()
    correo = request.form.get("correo", "").strip().lower()
    rol_id = request.form.get("rol_id", type=int)

    if not rol_id:
        flash("Debes seleccionar un rol.", "danger")
        return redirect(url_for("usuarios"))

    if str(current_user.id) == str(usuario_id):
        usuario_actual = obtener_registro("usuarios", usuario_id)
        if usuario_actual and int(usuario_actual.get("rol_id")) != int(rol_id):
            flash("No puedes cambiar el rol de tu propia sesion.", "danger")
            return redirect(url_for("usuarios"))

    try:
        actualizar_usuario(
            usuario_id=usuario_id,
            nombre=nombre,
            username=username,
            correo=correo,
            rol_id=rol_id,
        )
        flash("Usuario actualizado correctamente.", "success")
    except Exception as error:
        print(f"Error actualizando usuario: {error}")
        flash(str(error), "danger")

    return redirect(url_for("usuarios"))


@app.route(
    "/usuarios/<int:usuario_id>/password",
    methods=["POST"],
)
@login_required
@solo_administrador
def cambiar_password_web(usuario_id):
    password_nueva = request.form.get("password", "")

    try:
        cambiar_password(usuario_id, password_nueva)
        flash("Contrasena actualizada correctamente.", "success")
    except Exception as error:
        print(f"Error cambiando contrasena: {error}")
        flash(str(error), "danger")

    return redirect(url_for("usuarios"))


@app.route(
    "/usuarios/<int:usuario_id>/estado",
    methods=["POST"],
)
@login_required
@solo_administrador
def cambiar_estado_usuario(usuario_id):
    activo = request.form.get("activo") == "true"

    if str(current_user.id) == str(usuario_id) and not activo:
        flash("No puedes desactivar tu propio usuario.", "danger")
        return redirect(url_for("usuarios"))

    try:
        supabase.table("usuarios").update({
            "activo": activo
        }).eq("id", usuario_id).execute()
        limpiar_cache_usuarios(usuario_id)
        flash("Estado de usuario actualizado.", "success")
    except Exception as error:
        print(f"Error actualizando usuario: {error}")
        flash(str(error), "danger")

    return redirect(url_for("usuarios"))


@app.route("/usuarios/permisos", methods=["GET", "POST"])
@login_required
@solo_administrador
def administrar_permisos():
    roles = consultar_tabla("roles", "*", "nombre")
    modulos = consultar_tabla("modulos", "*", "orden")
    rol_id = request.values.get("rol_id", type=int)

    if request.method == "POST":
        if not rol_id:
            flash("Selecciona un rol.", "danger")
            return redirect(url_for("administrar_permisos"))

        rol = obtener_registro("roles", rol_id)
        if not rol:
            abort(404)

        if str(rol.get("nombre", "")).lower() == "administrador":
            flash(
                "Los permisos del Administrador siempre son totales.",
                "warning",
            )
            return redirect(
                url_for("administrar_permisos", rol_id=rol_id)
            )

        try:
            for modulo in modulos:
                modulo_id = modulo["id"]
                prefijo = f"modulo_{modulo_id}_"
                datos = {
                    "rol_id": rol_id,
                    "modulo_id": modulo_id,
                    "puede_ver": request.form.get(prefijo + "ver") == "on",
                    "puede_crear": request.form.get(prefijo + "crear") == "on",
                    "puede_editar": request.form.get(prefijo + "editar") == "on",
                    "puede_eliminar": request.form.get(prefijo + "eliminar") == "on",
                }

                existente = (
                    supabase
                    .table("permisos")
                    .select("id")
                    .eq("rol_id", rol_id)
                    .eq("modulo_id", modulo_id)
                    .limit(1)
                    .execute()
                )

                if existente.data:
                    (
                        supabase
                        .table("permisos")
                        .update(datos)
                        .eq("id", existente.data[0]["id"])
                        .execute()
                    )
                else:
                    supabase.table("permisos").insert(datos).execute()

            limpiar_cache_usuarios()
            flash("Permisos actualizados correctamente.", "success")
        except Exception as error:
            print(f"Error actualizando permisos: {error}")
            flash(str(error), "danger")

        return redirect(
            url_for("administrar_permisos", rol_id=rol_id)
        )

    permisos = []
    if rol_id:
        permisos = (
            supabase
            .table("permisos")
            .select("*")
            .eq("rol_id", rol_id)
            .execute()
        ).data or []

    permiso_por_modulo = {
        permiso["modulo_id"]: permiso
        for permiso in permisos
    }

    if plantilla_existe("usuarios/permisos.html"):
        return render_template(
            "usuarios/permisos.html",
            roles=roles,
            modulos=modulos,
            rol_id=rol_id,
            permiso_por_modulo=permiso_por_modulo,
        )

    return {
        "roles": roles,
        "modulos": modulos,
        "rol_id": rol_id,
        "permisos": permisos,
    }


# ==================================================
# ERRORES Y TEST
# ==================================================

@app.errorhandler(CSRFError)
def error_csrf(error):
    flash(
        "El formulario expiró o no es válido. Intenta de nuevo.",
        "danger",
    )
    return redirect(request.referrer or url_for("login"))


@app.errorhandler(403)
def acceso_denegado(error):
    if plantilla_existe("403.html"):
        return render_template("403.html"), 403
    return "No tienes permiso para realizar esta accion.", 403


@app.route("/test")
@login_required
def test():
    activos = consultar_tabla("activos")
    return {
        "conexion": True,
        "cantidad_activos": len(activos),
        "usuario": current_user.correo,
        "rol": current_user.rol_nombre,
    }


if __name__ == "__main__":
    app.run(
        host="0.0.0.0",
        port=5000,
        debug=True,
    )