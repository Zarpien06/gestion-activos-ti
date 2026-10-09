import io
from datetime import date, datetime, timedelta, timezone

from flask import send_file

MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
COLOR_ENCABEZADO = "1F3A5F"
COLOR_FILA_PAR = "F3F4F6"
MAX_FILAS_ZEBRA = 5000

FORMATOS = {
    "fecha": "dd/mm/yyyy",
    "fechahora": "dd/mm/yyyy hh:mm",
    "moneda": "#,##0;[Red]-#,##0",
    "entero": "#,##0",
    "decimal": "#,##0.00",
    "porcentaje": "0.0%",
}
NUMERICOS = {"moneda", "entero", "decimal", "porcentaje"}


def _ahora():
    # Hora de Colombia (UTC-5)
    return datetime.now(timezone.utc) - timedelta(hours=5)


def _nombre_hoja(libro, nombre):
    """Nombre valido (max 31, sin caracteres prohibidos) y no repetido."""
    limpio = "".join(c for c in str(nombre) if c not in '[]:*?/\\')[:31] or "Hoja"
    base, n = limpio, 2
    while limpio in libro.sheetnames:
        sufijo = f" {n}"
        limpio = base[: 31 - len(sufijo)] + sufijo
        n += 1
    return limpio


def _limpiar(valor):
    """None -> vacio; fechas con zona horaria -> sin zona (openpyxl no las admite)."""
    if valor is None:
        return ""
    if isinstance(valor, datetime) and valor.tzinfo is not None:
        return valor.replace(tzinfo=None)
    return valor


def _inferir_tipos(filas, n_cols):
    """Si la ruta no indica tipos, detecta fechas por el primer valor no vacio."""
    tipos = ["texto"] * n_cols
    for i in range(n_cols):
        for fila in filas[:50]:
            if i < len(fila) and fila[i] not in (None, ""):
                if isinstance(fila[i], datetime):
                    tipos[i] = "fechahora"
                elif isinstance(fila[i], date):
                    tipos[i] = "fecha"
                break
    return tipos


def _indice(encabezados, ref):
    """Acepta indice (0-based) o el texto del encabezado."""
    if isinstance(ref, int):
        return ref if 0 <= ref < len(encabezados) else None
    try:
        return list(encabezados).index(ref)
    except ValueError:
        return None


def _hoja_datos(libro, hoja):
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.properties import PageSetupProperties

    ws = libro.create_sheet(_nombre_hoja(libro, hoja["hoja"]))
    encabezados = list(hoja["encabezados"])
    filas = [list(f) for f in hoja["filas"]]
    n_cols = len(encabezados)

    tipos = list(hoja.get("tipos") or [])
    inferidos = _inferir_tipos(filas, n_cols)
    tipos = [(tipos[i] if i < len(tipos) and tipos[i] else inferidos[i]) for i in range(n_cols)]

    # Encabezado
    ws.append(encabezados)
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=COLOR_ENCABEZADO)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28

    # Datos
    anchos = [len(str(h)) for h in encabezados]
    for fila in filas:
        valores = [_limpiar(v) for v in fila]
        ws.append(valores)
        for i, valor in enumerate(valores[:n_cols]):
            largo = 10 if isinstance(valor, (date, datetime)) else len(str(valor))
            anchos[i] = max(anchos[i], largo)

    ultima = len(filas) + 1
    zebra = PatternFill("solid", fgColor=COLOR_FILA_PAR) if len(filas) <= MAX_FILAS_ZEBRA else None

    for r in range(2, ultima + 1):
        for i in range(n_cols):
            celda = ws.cell(row=r, column=i + 1)
            # Un texto que empieza con "=" no debe interpretarse como formula
            if isinstance(celda.value, str) and celda.value.startswith("="):
                celda.data_type = "s"
            formato = FORMATOS.get(tipos[i])
            if formato:
                celda.number_format = formato
            if tipos[i] in NUMERICOS:
                celda.alignment = Alignment(horizontal="right")
            if zebra and r % 2 == 1:
                celda.fill = zebra

    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(ancho, 50) + 2

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(n_cols)}{max(ultima, 1)}"

    # Semaforo de vencimientos (se recalcula al abrir el archivo)
    if filas:
        rojo = PatternFill("solid", bgColor="FEE2E2")
        ambar = PatternFill("solid", bgColor="FEF3C7")
        for ref in hoja.get("vence", []):
            idx = _indice(encabezados, ref)
            if idx is None:
                continue
            L = get_column_letter(idx + 1)
            rango = f"{L}2:{L}{ultima}"
            ws.conditional_formatting.add(rango, FormulaRule(
                formula=[f'AND(ISNUMBER({L}2),{L}2<TODAY())'],
                fill=rojo, font=Font(color="B91C1C", bold=True)))
            ws.conditional_formatting.add(rango, FormulaRule(
                formula=[f'AND(ISNUMBER({L}2),{L}2>=TODAY(),{L}2<=TODAY()+90)'],
                fill=ambar, font=Font(color="B45309")))

    # Totales con SUBTOTAL: cambian cuando el usuario filtra en Excel
    totales = [i for i in (_indice(encabezados, t) for t in hoja.get("totales", [])) if i is not None]
    if totales and filas:
        fila_t = ultima + 2
        etiqueta = ws.cell(row=fila_t, column=1, value="TOTAL")
        etiqueta.font = Font(bold=True)
        borde = Border(top=Side(style="thin", color="9CA3AF"))
        for i in totales:
            L = get_column_letter(i + 1)
            c = ws.cell(row=fila_t, column=i + 1, value=f"=SUBTOTAL(109,{L}2:{L}{ultima})")
            c.font = Font(bold=True)
            c.number_format = FORMATOS.get(tipos[i], "#,##0")
            c.alignment = Alignment(horizontal="right")
            c.border = borde

    # Impresion
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = "1:1"


def _hoja_resumen(libro, bloques, subtitulo, titulo, filtros, kpis):
    from openpyxl.chart import BarChart, PieChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.styles import Alignment, Font, PatternFill

    ws = libro.create_sheet("Resumen", 0)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 3

    ws["A1"] = titulo
    ws["A1"].font = Font(bold=True, size=16, color=COLOR_ENCABEZADO)
    ws["A2"] = subtitulo
    ws["A2"].font = Font(italic=True, color="6B7280")

    fila = 4

    # Filtros aplicados
    ws.cell(fila, 1, "Filtros aplicados").font = Font(bold=True, size=12)
    fila += 1
    activos = {k: v for k, v in (filtros or {}).items() if v not in (None, "")}
    if activos:
        for k, v in activos.items():
            ws.cell(fila, 1, k).font = Font(color="6B7280")
            ws.cell(fila, 2, str(v))
            fila += 1
    else:
        ws.cell(fila, 1, "Sin filtros (todos los registros)").font = Font(color="6B7280")
        fila += 1

    # Indicadores
    if kpis:
        fila += 1
        ws.cell(fila, 1, "Indicadores").font = Font(bold=True, size=12)
        fila += 1
        for kpi in kpis:
            etiqueta, valor = kpi[0], kpi[1]
            tipo = kpi[2] if len(kpi) > 2 else None
            ws.cell(fila, 1, etiqueta)
            c = ws.cell(fila, 2, valor)
            c.font = Font(bold=True, size=12)
            c.alignment = Alignment(horizontal="right")
            if tipo in FORMATOS:
                c.number_format = FORMATOS[tipo]
            fila += 1

    fila += 1
    for bloque in bloques:
        datos = bloque.get("datos") or []
        n = len(datos)

        ws.cell(fila, 1, bloque["titulo"]).font = Font(bold=True, size=12)
        for col, texto in ((1, "Concepto"), (2, "Total")):
            c = ws.cell(fila + 1, col, texto)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=COLOR_ENCABEZADO)

        if not datos:
            ws.cell(fila + 2, 1, "Sin datos").font = Font(italic=True, color="9CA3AF")
            fila += 5
            continue

        for i, (etiqueta, valor) in enumerate(datos):
            ws.cell(fila + 2 + i, 1, "" if etiqueta is None else etiqueta)
            c = ws.cell(fila + 2 + i, 2, valor)
            if bloque.get("formato") in FORMATOS:
                c.number_format = FORMATOS[bloque["formato"]]

        tipo = bloque.get("grafico")
        if tipo:
            if tipo == "pie" and n <= 8:
                grafico = PieChart()
                grafico.dataLabels = DataLabelList()
                grafico.dataLabels.showPercent = True
            else:
                grafico = BarChart()
                grafico.type = "col"
                grafico.legend = None

            grafico.title = bloque["titulo"]
            grafico.height = 7.5
            grafico.width = 15
            grafico.add_data(
                Reference(ws, min_col=2, min_row=fila + 1, max_row=fila + 1 + n),
                titles_from_data=True,
            )
            grafico.set_categories(
                Reference(ws, min_col=1, min_row=fila + 2, max_row=fila + 1 + n)
            )
            ws.add_chart(grafico, f"D{fila}")
            fila += max(n + 4, 18)
        else:
            fila += n + 4


def respuesta_libro(nombre_archivo, hojas, resumen=None, titulo="Resumen estadistico",
                    filtros=None, kpis=None, usuario=""):
    """Libro .xlsx con una hoja de resumen (graficas) y una o varias hojas de datos.

    hojas:   [{"hoja": str, "encabezados": [...], "filas": [[...], ...],
               # opcionales:
               "tipos": ["texto", "fecha", "moneda", "entero", ...],
               "totales": ["Costo", 5],        # encabezado o indice
               "vence": ["Fecha de fin"]}, ...]
    resumen: [{"titulo": str, "datos": [(etiqueta, valor), ...],
               "grafico": "bar" | "pie" | None, "formato": "moneda" (opcional)}, ...]
    filtros: {"Estado": "Vencidas", ...}  (solo se muestran los que tengan valor)
    kpis:    [("Licencias", 25, "entero"), ("Costo total", 1200000, "moneda")]
    usuario: nombre de quien genera el reporte
    """
    from openpyxl import Workbook

    libro = Workbook()
    libro.remove(libro.active)

    for hoja in hojas:
        _hoja_datos(libro, hoja)

    if resumen or kpis or filtros:
        subtitulo = f"Generado el {_ahora():%Y-%m-%d %H:%M}"
        if usuario:
            subtitulo += f"  ·  Usuario: {usuario}"
        _hoja_resumen(libro, resumen or [], subtitulo, titulo, filtros, kpis)
        libro.active = 0

    buffer = io.BytesIO()
    libro.save(buffer)
    buffer.seek(0)

    return send_file(
        buffer,
        mimetype=MIME_XLSX,
        as_attachment=True,
        download_name=nombre_archivo,
    )


def respuesta_excel(nombre_archivo, hoja, encabezados, filas):
    """Compatibilidad: un Excel simple de una sola hoja."""
    return respuesta_libro(
        nombre_archivo,
        [{"hoja": hoja, "encabezados": encabezados, "filas": filas}],
    )